"""
文件解析模块
从各类文件中提取纯文本，供 Circle 读取上传文档后回答。
支持 18 种格式：txt / md / pdf / docx / rtf / html / xlsx / xls / csv /
json / xml / yaml / yml / pptx / jpg / jpeg / png / gif（图片含可选 OCR）。
可选解析库按需导入，缺失时抛 ImportError，不影响其他功能。
"""

import json
import os

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5MB
MAX_TEXT_LEN = 8000  # 截断长度
ALLOWED_EXTS = {
    # 文档：txt, pdf, docx, rtf, html
    "txt",
    "pdf",
    "docx",
    "rtf",
    "html",
    "htm",
    "md",
    # 表格：xlsx, xls, csv
    "xlsx",
    "xls",
    "csv",
    # 数据：json, xml, yaml, yml
    "json",
    "xml",
    "yaml",
    "yml",
    # 演示：pptx
    "pptx",
    # 图片：jpg, jpeg, png, gif
    "jpg",
    "jpeg",
    "png",
    "gif",
}


# ── 文件解析库（按需导入，缺失不影响其他功能） ──
def _import_mod(m: str, *names: str):
    """按需导入模块或其中的符号，缺失时由调用方捕获。"""
    return __import__(m, fromlist=list(names)) if names else __import__(m)


_HAS = {}
for _lib, _mods in [
    ("pypdf", [("PdfReader", "pypdf")]),
    ("docx", [("Document", "docx")]),
    ("openpyxl", [("load_workbook", "openpyxl")]),
    ("xlrd", [("open_workbook", "xlrd")]),
    ("striprtf", [("rtf_to_text", "striprtf")]),
    ("bs4", [("BeautifulSoup", "bs4")]),
    ("yaml", [("safe_load", "yaml")]),
    ("pptx", [("Presentation", "pptx")]),
    ("PIL", [("Image", "PIL")]),
]:
    try:
        _mod = _import_mod(_lib)
        for _attr, _pkg in _mods:
            _HAS[_pkg] = True
    except ImportError:
        for _, _pkg in _mods:
            _HAS[_pkg] = False

# OCR 单独处理（依赖 tesseract 系统安装）
try:
    import pytesseract

    HAS_TESSERACT = True
except ImportError:
    HAS_TESSERACT = False


def extract_text(file_path: str) -> str:
    """根据扩展名自动选择解析库提取文本，返回纯文本。"""
    ext = os.path.splitext(file_path)[1].lower().lstrip(".")

    # ── 纯文本类 ──
    if ext in ("txt", "md"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    # ── PDF ──
    if ext == "pdf":
        if not _HAS.get("pypdf"):
            raise ImportError("缺少 pypdf 库，请 pip install pypdf")
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    # ── Word ──
    if ext == "docx":
        if not _HAS.get("docx"):
            raise ImportError("缺少 python-docx 库，请 pip install python-docx")
        from docx import Document

        return "\n".join(p.text for p in Document(file_path).paragraphs)

    # ── RTF ──
    if ext == "rtf":
        if not _HAS.get("striprtf"):
            raise ImportError("缺少 striprtf 库，请 pip install striprtf")
        from striprtf.striprtf import rtf_to_text

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return rtf_to_text(f.read())

    # ── HTML ──
    if ext in ("html", "htm"):
        if not _HAS.get("bs4"):
            raise ImportError("缺少 beautifulsoup4 库，请 pip install beautifulsoup4 lxml")
        from bs4 import BeautifulSoup

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return BeautifulSoup(f.read(), "lxml").get_text("\n", strip=True)

    # ── Excel (.xlsx) ──
    if ext == "xlsx":
        if not _HAS.get("openpyxl"):
            raise ImportError("缺少 openpyxl 库，请 pip install openpyxl")
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        lines = []
        for name in wb.sheetnames:
            ws = wb[name]
            lines.append(f"[Sheet: {name}]")
            for row in ws.iter_rows(values_only=True):
                line = "\t".join(str(c) if c is not None else "" for c in row)
                if line.strip():
                    lines.append(line)
        wb.close()
        return "\n".join(lines)

    # ── Excel (.xls 旧版) ──
    if ext == "xls":
        if not _HAS.get("xlrd"):
            raise ImportError("缺少 xlrd 库，请 pip install xlrd")
        import xlrd

        wb = xlrd.open_workbook(file_path)
        lines = []
        for name in wb.sheet_names():
            ws = wb.sheet_by_name(name)
            lines.append(f"[Sheet: {name}]")
            for r in range(ws.nrows):
                line = "\t".join(str(ws.cell_value(r, c)) for c in range(ws.ncols))
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)

    # ── CSV ──
    if ext == "csv":
        import csv

        with open(file_path, "r", encoding="utf-8-sig", errors="ignore") as f:
            reader = csv.reader(f)
            return "\n".join("\t".join(row) for row in reader if any(row))

    # ── JSON ──
    if ext == "json":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            data = json.load(f)
        return json.dumps(data, ensure_ascii=False, indent=2)

    # ── XML ──
    if ext == "xml":
        import xml.etree.ElementTree as ET

        tree = ET.parse(file_path)

        # 递归提取所有文本
        def _walk(elem, depth=0):
            texts = []
            if elem.text and elem.text.strip():
                texts.append("  " * depth + elem.text.strip())
            for child in elem:
                texts.extend(_walk(child, depth + 1))
                if child.tail and child.tail.strip():
                    texts.append("  " * depth + child.tail.strip())
            return texts

        return "\n".join(_walk(tree.getroot()))

    # ── YAML ──
    if ext in ("yaml", "yml"):
        if not _HAS.get("yaml"):
            raise ImportError("缺少 pyyaml 库，请 pip install pyyaml")
        import yaml

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            data = yaml.safe_load(f)
        return yaml.dump(data, allow_unicode=True, default_flow_style=False)

    # ── PowerPoint ──
    if ext == "pptx":
        if not _HAS.get("pptx"):
            raise ImportError("缺少 python-pptx 库，请 pip install python-pptx")
        from pptx import Presentation

        prs = Presentation(file_path)
        lines = []
        for i, slide in enumerate(prs.slides, 1):
            lines.append(f"[Slide {i}]")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            lines.append(t)
        return "\n".join(lines)

    # ── 图片（基本信息 + 可选 OCR） ──
    if ext in ("jpg", "jpeg", "png", "gif"):
        if not _HAS.get("PIL"):
            raise ImportError("缺少 Pillow 库，请 pip install pillow")
        from PIL import Image

        img = Image.open(file_path)
        info = f"[图片信息] 格式={img.format}  尺寸={img.size[0]}x{img.size[1]}  模式={img.mode}"
        # 尝试 OCR
        if HAS_TESSERACT:
            try:
                text = pytesseract.image_to_string(img, lang="chi_sim+eng")
                if text.strip():
                    return info + "\n[OCR 识别文字]\n" + text.strip()
            except Exception:
                pass
        return info

    raise ValueError(f"不支持的文件类型或缺少解析库：{ext}")
